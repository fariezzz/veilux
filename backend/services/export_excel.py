"""Layanan ekspor laporan benchmark evaluasi forensik ke format Microsoft Excel (.xlsx).

Menggunakan openpyxl untuk menghasilkan berkas spreadsheet terformat rapi:
- Lebar kolom proporsional (auto-width dengan padding)
- Wrap text pada seluruh sel data
- Pemformatan angka metrik presisi (PSNR, MSE, NC, BER, Tamper Ratio)
- Pewarnaan status deteksi watermark (hijau untuk terdeteksi, merah untuk rusak)
"""

from __future__ import annotations

from datetime import datetime
import io
import math
from typing import Any, Dict, List, Optional, Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generate_benchmark_xlsx(
    results: List[Dict[str, Any]],
    metadata: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Menghasilkan berkas Excel (.xlsx) dari data benchmark serangan citra.

    Args:
        results: Daftar baris hasil benchmark per serangan.
        metadata: Informasi tambahan seperti tipe watermark, tanggal uji, dll.

    Returns:
        bytes: Konten biner berkas XLSX siap diunduh atau dikirim via HTTP.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Hasil Benchmark"

    # Pastikan grid lines selalu aktif saat dibuka di Excel
    ws.views.sheetView[0].showGridLines = True

    # Definisi Palet Warna dan Gaya
    font_title = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    font_subtitle = Font(name="Segoe UI", size=9, italic=True, color="475569")
    font_meta_label = Font(name="Segoe UI", size=9, bold=True, color="334155")
    font_meta_value = Font(name="Segoe UI", size=9, color="1E293B")

    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    font_data = Font(name="Segoe UI", size=9, color="0F172A")
    font_data_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")

    fill_ok = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_ok = Font(name="Segoe UI", size=9, bold=True, color="166534")

    fill_fail = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_fail = Font(name="Segoe UI", size=9, bold=True, color="991B1B")

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    data_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    header_border = Border(
        left=Side(border_style="thin", color="475569"),
        right=Side(border_style="thin", color="475569"),
        top=Side(border_style="thin", color="475569"),
        bottom=Side(border_style="medium", color="0F172A"),
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # 1. Judul dan Metadata Header Laporan
    ws.merge_cells("A1:J1")
    ws["A1"] = "VEILUX - LAPORAN EVALUASI BENCHMARK KETAHANAN WATERMARK"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:J2")
    ws["A2"] = "Pengujian Fragile Least Significant Bit (LSB v3) & Lokalisasi Manipulasi Blok terhadap 8 Jenis Serangan Citra"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    # Metadata informasi uji
    meta = metadata or {}
    export_time = meta.get("export_time") or datetime.now().strftime("%d-%m-%Y %H:%M:%S WIB")
    wm_type = meta.get("watermark_type", "TEXT / LOGO").upper()

    ws["A4"] = "Waktu Pengujian:"
    ws["A4"].font = font_meta_label
    ws["B4"] = export_time
    ws["B4"].font = font_meta_value

    ws["D4"] = "Tipe Watermark:"
    ws["D4"].font = font_meta_label
    ws["E4"] = wm_type
    ws["E4"].font = font_meta_value

    ws["G4"] = "Total Serangan:"
    ws["G4"].font = font_meta_label
    ws["H4"] = f"{len(results)} Skenario"
    ws["H4"].font = font_meta_value

    # 2. Header Tabel Data
    headers = [
        ("No", align_center, 6),
        ("Jenis Serangan", align_left, 24),
        ("Kategori Serangan", align_left, 22),
        ("PSNR (dB)", align_right, 14),
        ("MSE", align_right, 14),
        ("NC", align_right, 14),
        ("BER", align_right, 14),
        ("Status Watermark", align_center, 18),
        ("Blok Valid", align_center, 14),
        ("Tamper Ratio", align_right, 15),
    ]

    header_row_idx = 6
    ws.row_dimensions[header_row_idx].height = 28

    for col_idx, (col_name, col_align, _) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = col_align
        cell.border = header_border

    # 3. Baris Data Benchmark
    current_row = header_row_idx + 1
    for idx, row in enumerate(results):
        ws.row_dimensions[current_row].height = 22
        is_even = idx % 2 == 1

        # Parsing nilai numerik dengan aman
        no_val = row.get("no", idx + 1)
        name_val = str(row.get("name", ""))
        cat_val = str(row.get("category", ""))

        def _to_float(v: Any) -> Optional[float]:
            if v is None:
                return None
            if isinstance(v, (int, float)):
                return float(v)
            try:
                s = str(v).replace("%", "").replace(",", ".").strip()
                if s in ("--", "N/A", "N/A (Crop)", "null", "None", ""):
                    return None
                return float(s)
            except (ValueError, TypeError):
                return None

        psnr_num = _to_float(row.get("psnr"))
        mse_num = _to_float(row.get("mse"))
        nc_num = _to_float(row.get("nc"))
        ber_num = _to_float(row.get("ber"))

        # Tamper ratio
        tr_raw = row.get("tamperRatio") or row.get("tamper_ratio")
        tr_num = _to_float(tr_raw)
        if tr_num is not None and tr_num > 1.0 and "%" in str(tr_raw):
            tr_num = tr_num / 100.0  # Konversi persen (contoh: 85.5% -> 0.855)

        detected_flag = bool(row.get("detected", False))
        valid_blocks_str = str(row.get("validBlocks") or row.get("valid_blocks", "--"))

        # Pengisian sel data
        c1 = ws.cell(row=current_row, column=1, value=no_val)
        c1.alignment = align_center
        c1.number_format = "0"

        c2 = ws.cell(row=current_row, column=2, value=name_val)
        c2.alignment = align_left
        c2.font = font_data_bold

        c3 = ws.cell(row=current_row, column=3, value=cat_val)
        c3.alignment = align_left

        c4 = ws.cell(row=current_row, column=4, value=psnr_num if psnr_num is not None else str(row.get("psnr", "--")))
        c4.alignment = align_right
        if psnr_num is not None:
            c4.number_format = "0.00"

        c5 = ws.cell(row=current_row, column=5, value=mse_num if mse_num is not None else str(row.get("mse", "--")))
        c5.alignment = align_right
        if mse_num is not None:
            c5.number_format = "0.0000"

        c6 = ws.cell(row=current_row, column=6, value=nc_num if nc_num is not None else str(row.get("nc", "--")))
        c6.alignment = align_right
        if nc_num is not None:
            c6.number_format = "0.0000"

        c7 = ws.cell(row=current_row, column=7, value=ber_num if ber_num is not None else str(row.get("ber", "--")))
        c7.alignment = align_right
        if ber_num is not None:
            c7.number_format = "0.0000"

        status_label = "Terdeteksi" if detected_flag else "Rusak"
        c8 = ws.cell(row=current_row, column=8, value=status_label)
        c8.alignment = align_center
        c8.font = font_ok if detected_flag else font_fail
        c8.fill = fill_ok if detected_flag else fill_fail

        c9 = ws.cell(row=current_row, column=9, value=valid_blocks_str)
        c9.alignment = align_center

        c10 = ws.cell(row=current_row, column=10, value=tr_num if tr_num is not None else str(tr_raw or "--"))
        c10.alignment = align_right
        if tr_num is not None:
            c10.number_format = "0.0%"

        # Border dan background selang-seling (zebra)
        for col_idx in range(1, 11):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = data_border
            if col_idx != 8:
                cell.font = font_data if col_idx != 2 else font_data_bold
                if is_even:
                    cell.fill = fill_zebra

        current_row += 1

    # 4. Ringkasan Evaluasi / Rata-rata Metrik (Footer Table)
    summary_row = current_row
    ws.row_dimensions[summary_row].height = 24
    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=3)
    sum_label = ws.cell(row=summary_row, column=1, value="Rata-rata Evaluasi:")
    sum_label.font = font_header
    sum_label.fill = fill_header
    sum_label.alignment = Alignment(horizontal="right", vertical="center")

    # Formula rata-rata untuk PSNR, MSE, NC, BER
    data_start = header_row_idx + 1
    data_end = current_row - 1

    sum_psnr = ws.cell(row=summary_row, column=4, value=f"=AVERAGE(D{data_start}:D{data_end})")
    sum_psnr.number_format = "0.00"

    sum_mse = ws.cell(row=summary_row, column=5, value=f"=AVERAGE(E{data_start}:E{data_end})")
    sum_mse.number_format = "0.0000"

    sum_nc = ws.cell(row=summary_row, column=6, value=f"=AVERAGE(F{data_start}:F{data_end})")
    sum_nc.number_format = "0.0000"

    sum_ber = ws.cell(row=summary_row, column=7, value=f"=AVERAGE(G{data_start}:G{data_end})")
    sum_ber.number_format = "0.0000"

    # Kosongkan kolom 8-9 untuk footer
    ws.cell(row=summary_row, column=8, value="")
    ws.cell(row=summary_row, column=9, value="")

    sum_tr = ws.cell(row=summary_row, column=10, value=f"=AVERAGE(J{data_start}:J{data_end})")
    sum_tr.number_format = "0.0%"

    for col_idx in range(1, 11):
        c = ws.cell(row=summary_row, column=col_idx)
        c.font = font_header
        c.fill = fill_header
        c.border = header_border
        if col_idx in (4, 5, 6, 7, 10):
            c.alignment = align_right

    # 5. Penyesuaian Lebar Kolom (Auto-fit dengan Padding)
    for col_idx, (_, _, min_w) in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = min_w
        for row_idx in range(header_row_idx, summary_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)) + 3)
        ws.column_dimensions[col_letter].width = min(max_len, 35)

    # Freeze panes di bawah header tabel
    ws.freeze_panes = f"A{header_row_idx + 1}"

    # Simpan ke buffer byte in-memory
    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream.getvalue()
